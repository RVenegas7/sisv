#!/usr/bin/env python3
"""Carga CRUDA en PostgreSQL del archivo ``lara_2019_2020.xlsx``.

El archivo (raiz del repo, ignorado por git) son las defunciones de Lara de
2019-2020 persona a persona, en 220 columnas (estructura EV-14), con hoja
``Diccionario``. **No es una tabla de Oracle**: es un Excel suelto, y por eso
NO entra al esquema ``sismai`` (que debe reflejar Oracle tal cual). Se carga
en un esquema aparte, ``espejo_extra``, como tabla cruda: todas las columnas
en TEXT, mas ``_hoja`` y ``_fila`` de procedencia, indice por ``planilla``.

La clave de enlace con el legacy es la columna ``PLANILLA`` del xlsx =
``sismai."CERTIFICADO"."NUMEROMSDS"`` (verificado 09/10/2026: 4.803 de los
5.729 certificados de 2019-20 aparecen en el xlsx).

Uso:
    python migracion/cargar_lara_2019_2020_pg.py [--archivo lara_2019_2020.xlsx]
        [--esquema espejo_extra] [--tabla lara_2019_2020] [--ejecutar]

Sin ``--ejecutar`` solo informa (dry-run). Idempotente: DROP + CREATE + COPY.

Solo lee el Excel con ``openpyxl`` y escribe en PG; no toca ``registros`` ni
el ETL. Conciliar estas filas en ``registros.Defuncion`` es una decision
aparte (atribucion por centro y dedup), documentada en PENDIENTES.
"""
import argparse
import csv
import datetime
import io
import os
import re
import sys
from pathlib import Path

import openpyxl
import psycopg

DB = {
    "host": os.environ.get("PGHOST", "localhost"),
    "port": os.environ.get("PGPORT", "5436"),
    "dbname": os.environ.get("PGDATABASE", "sis_salud_db"),
    "user": os.environ.get("PGUSER", "sis_user"),
    "password": os.environ.get("PGPASSWORD", "sis_password"),
}

HOJAS = ("2019", "2020")  # la hoja "Diccionario" no es dato


def nombre_columna(crudo, usados):
    """Normaliza un rotulo de la fila 2 a un identificador SQL seguro y unico."""
    base = re.sub(r"[^0-9a-zA-Z]+", "_", str(crudo).strip()).strip("_").lower()
    if not base:
        base = "col"
    if base[0].isdigit():
        base = "c_" + base
    if len(base) > 60:
        base = base[:60]
    nombre, n = base, 2
    while nombre in usados:
        nombre = f"{base}_{n}"
        n += 1
    usados.add(nombre)
    return nombre


def a_texto(v):
    """Convierte una celda a texto o None (None -> NULL en el COPY)."""
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat(sep=" ") if isinstance(v, datetime.datetime) else v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def leer(archivo):
    """Devuelve (encabezados, filas) de las hojas de datos."""
    wb = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    usados = set()
    encabezados = None
    filas = []
    for hoja in HOJAS:
        if hoja not in wb.sheetnames:
            continue
        ws = wb[hoja]
        it = ws.iter_rows(min_row=2, max_row=2, values_only=True)
        cab = list(next(it))
        if encabezados is None:
            encabezados = [nombre_columna(c, usados) for c in cab]
            ncols = len(encabezados)
            idx_planilla = cab.index("PLANILLA") if "PLANILLA" in cab else None
        for n, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
            r = list(r)[:ncols]
            r += [None] * (ncols - len(r))
            valores = [a_texto(v) for v in r]
            planilla = valores[idx_planilla] if idx_planilla is not None else None
            filas.append([hoja, n, planilla] + valores)
    return encabezados, filas


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--archivo", default="lara_2019_2020.xlsx")
    ap.add_argument("--esquema", default="espejo_extra")
    ap.add_argument("--tabla", default="lara_2019_2020")
    ap.add_argument("--ejecutar", action="store_true",
                    help="aplica los cambios (sin esto es dry-run)")
    args = ap.parse_args()

    archivo = Path(args.archivo)
    if not archivo.is_file():
        sys.exit(f"No existe el archivo {archivo}")

    print(f"Leyendo {archivo} ...", flush=True)
    encabezados, filas = leer(archivo)
    if not encabezados:
        sys.exit("No se encontraron hojas de datos 2019/2020")
    cols = ["_hoja", "_fila", "_planilla"] + encabezados
    por_hoja = {}
    for f in filas:
        por_hoja[f[0]] = por_hoja.get(f[0], 0) + 1
    print(f"  {len(filas)} filas ({por_hoja}) x {len(cols)} columnas")
    print(f"  destino: {args.esquema}.{args.tabla}")

    if not args.ejecutar:
        print("dry-run: usa --ejecutar para cargar")
        return

    definicion = ",\n  ".join(f'"{c}" text' for c in cols)
    lista_cols = ", ".join(f'"{c}"' for c in cols)
    with psycopg.connect(**DB) as con:
        with con.cursor() as cur:
            cur.execute(f'CREATE SCHEMA IF NOT EXISTS "{args.esquema}"')
            cur.execute(f'DROP TABLE IF EXISTS "{args.esquema}"."{args.tabla}"')
            cur.execute(
                f'CREATE TABLE "{args.esquema}"."{args.tabla}" (\n  {definicion}\n)'
            )
            copy_sql = (
                f'COPY "{args.esquema}"."{args.tabla}" ({lista_cols}) '
                "FROM STDIN WITH (FORMAT csv)"
            )
            with cur.copy(copy_sql) as cp:
                buf = io.StringIO()
                w = csv.writer(buf, lineterminator="\n")
                for f in filas:
                    w.writerow(f)
                    if buf.tell() > 1 << 20:
                        cp.write(buf.getvalue())
                        buf.seek(0)
                        buf.truncate(0)
                cp.write(buf.getvalue())
            cur.execute(
                f'CREATE INDEX ix_{args.tabla}_planilla '
                f'ON "{args.esquema}"."{args.tabla}" (_planilla)'
            )
        con.commit()
    print(f"Cargadas {len(filas)} filas en {args.esquema}.{args.tabla}")


if __name__ == "__main__":
    main()
